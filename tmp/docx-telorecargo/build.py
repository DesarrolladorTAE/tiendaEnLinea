from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

root = Path('C:/Users/resid/Documents/MTELMX')
reference = root / 'docs/manual-usuario-productos.docx'
source = (root / 'docs/tutorial-recargas-telorecargo.md').read_text(encoding='utf-8')
doc = Document(reference)
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
for name in ['Title', 'Heading 1']:
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
if 'Heading 2' not in doc.styles:
    style = doc.styles.add_style('Heading 2', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Normal']
    style.font.size = Pt(12)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.widow_control = True

def inline(paragraph, text):
    for piece in re.split(r'(\*\*.*?\*\*)', text):
        if not piece:
            continue
        run = paragraph.add_run(piece[2:-2] if piece.startswith('**') else piece)
        if piece.startswith('**'):
            run.bold = True

def para(text, style=None):
    p = doc.add_paragraph(style=style)
    inline(p, text)
    return p

para('Manual de usuario de recargas Telorecargo', 'Title')
para('Mi Tienda en Línea MX · Versión 1.0 · 29 de octubre de 2026')
para('Guía para personal de caja y administración. Explica cómo vender recargas, registrar el cobro, entregar tickets, consultar operaciones y solicitar saldo.')

lines = source.splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line or line.startswith('# ') or line == 'Guía para personal de caja y administración de tienda.':
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([cell.strip() for cell in lines[i].strip().strip('|').split('|')])
            i += 1
        rows = [rows[0]] + rows[2:]
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.autofit = False
        widths = [Cm(5.2), Cm(11.8)]
        for col, width in zip(table.columns, widths):
            col.width = width
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            edge_el = OxmlElement('w:' + edge)
            for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
                edge_el.set(qn('w:' + key), value)
            borders.append(edge_el)
        props.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for edge in ['top', 'left', 'bottom', 'right']:
            el = OxmlElement('w:' + edge)
            el.set(qn('w:w'), '100')
            el.set(qn('w:type'), 'dxa')
            margins.append(el)
        props.append(margins)
        for index, values in enumerate(rows):
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell, width, value in zip(row.cells, widths, values):
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.keep_with_next = index < len(rows) - 1
                inline(p, value)
                for run in p.runs:
                    run.font.size = Pt(10)
                    if index == 0:
                        run.bold = True
                if index == 0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'E9EFF5')
                    cell._tc.get_or_add_tcPr().append(shading)
        para('')
        continue
    if line.startswith('## '):
        para(line[3:].replace('.', '', 1), 'Heading 1')
    elif line.startswith('### '):
        para(line[4:], 'Heading 2')
    else:
        p = para(line)
        if re.match(r'^\d+\. ', line):
            p.paragraph_format.left_indent = Cm(.5)
            p.paragraph_format.first_line_indent = Cm(-.5)
            previous = p._p.getprevious()
            if previous is not None and previous.tag == qn('w:p'):
                from docx.text.paragraph import Paragraph
                Paragraph(previous, doc).paragraph_format.keep_with_next = True
    i += 1

doc.core_properties.title = 'Manual de usuario de recargas Telorecargo'
doc.core_properties.subject = 'Venta de recargas y entrega de comprobantes en el POS'
doc.core_properties.author = 'Mi Tienda en Línea MX'
doc.core_properties.last_modified_by = 'Mi Tienda en Línea MX'
output = root / 'docs/manual-usuario-recargas-telorecargo.docx'
doc.save(output)
(root / 'tmp/docx-telorecargo/artifact.md').write_text(
    '# Referencia de formato\n\n'
    f'Referencia: {reference}\nSHA256: {sha256(reference.read_bytes()).hexdigest()}\n\n'
    'Se reutiliza el paquete del manual de Productos, su sección A4, márgenes de 2 cm, '
    'Calibri 11, título de 20 puntos, Heading 1 de 14 puntos y espaciado de párrafos. '
    'El cuerpo se sustituye por el tutorial completo. Se omiten tablas de captura vacías '
    'porque no forman parte del contenido del tutorial. Los títulos se ajustan a negro. '
    'Se agregan subtítulos de 12 puntos y tablas de datos con anchos 5.2 y 11.8 cm, '
    'bordes grises, filas expandibles y encabezados repetidos. '
    'El renderer empaquetado falla porque no existe soffice en este runtime Windows; '
    'se utiliza la exportación nativa de Word para revisar páginas.\n', encoding='utf-8')
print(output)
